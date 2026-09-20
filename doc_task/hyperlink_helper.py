from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

def add_hyperlink(paragraph, url: str, text: str):
    """
    Add a clickable hyperlink to a python-docx paragraph.
    
    :param paragraph: The paragraph to add the hyperlink to.
    :param url: The URL for the hyperlink.
    :param text: The display text for the hyperlink.
    """
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0000FF')  # Blue color
    rPr.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')  # Single underline
    rPr.append(underline)
    
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    r.append(t)
    
    hyperlink.append(r)
    paragraph._p.append(hyperlink)

def append_yt_links_to_doc(app_data, urls: list, original_query: str):
    """
    Append YouTube links as bullets under a Heading 2 in the active DOCX document.
    """
    file = app_data.get_file_name()
    folder = app_data.get_folder_path()
    path = f"{folder}/{file}"
    if not os.path.isfile(path):
        raise FileNotFoundError(f"Document not found at: {path}")

    doc = Document(path)
    p = doc.add_paragraph(f"YouTube Suggestions for '{original_query}':", style="Heading 2")
    p.paragraph_format.space_after = 0
    for url in urls:
        p = doc.add_paragraph("", style="List Bullet")
        add_hyperlink(p, url, url)
        p.paragraph_format.left_indent = Inches(0.5)
    doc.save(path)
