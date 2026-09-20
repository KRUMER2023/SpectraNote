from docx import Document
from docx.shared import Inches, Pt
import os


def append_summary_to_doc(app_data, summary_text: str, original_preview: str = ""):
    """
    Append an AI/T5 generated abstractive summary into the active DOCX document.

    Format:
    Summary: (Bold Normal Paragraph)
    Generated summary text with proper paragraph formatting...
    """
    file = app_data.get_file_name()
    folder = app_data.get_folder_path()
    path = f"{folder}/{file}"
    if not os.path.isfile(path):
        raise FileNotFoundError(f"Document not found at: {path}")

    doc = Document(path)

    # 1. Header (Bold Normal Paragraph)
    p_header = doc.add_paragraph(style="Normal")
    p_header.paragraph_format.space_before = Pt(6)
    p_header.paragraph_format.space_after = Pt(2)
    run_header = p_header.add_run("Summary:")
    run_header.bold = True

    # 2. Summary Body
    p_body = doc.add_paragraph(style="Normal")
    p_body.paragraph_format.left_indent = Inches(0.2)
    p_body.paragraph_format.space_before = Pt(2)
    p_body.paragraph_format.space_after = Pt(4)
    p_body.add_run(summary_text.strip())

    doc.save(path)
