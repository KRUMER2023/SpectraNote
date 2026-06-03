from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import requests
import os
import tkinter as tk
from tkinter import messagebox, Checkbutton, IntVar, Button, Label, Frame, Canvas, Scrollbar

API_KEY = os.getenv('YOUTUBE_API_KEY')
BASE_URL = "https://www.googleapis.com/youtube/v3/search"

def add_hyperlink(paragraph, url, text):
    """
    Add a hyperlink to a paragraph.
    
    :param paragraph: The paragraph to add the hyperlink to.
    :param url: The URL for the hyperlink.
    :param text: The display text for the hyperlink.
    """
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0000FF')  # Blue color
    rPr.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')  # Single underline
    rPr.append(underline)
    
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    r.append(t)
    
    hyperlink.append(r)
    paragraph._p.append(hyperlink)

def fetch_yt_suggestions(query: str, max_results: int = 4):
    """Fetch top 4 YouTube video suggestions for the query."""
    if not API_KEY:
        raise ValueError("YOUTUBE_API_KEY environment variable is not set. Please check your .env file.")
    params = {
        'part': 'snippet',
        'q': query,
        'type': 'video',
        'order': 'relevance',
        'maxResults': max_results,
        'key': API_KEY
    }
    response = requests.get(BASE_URL, params=params)
    if response.status_code != 200:
        raise ValueError(f"API error: {response.text}")
    data = response.json()
    if 'error' in data:
        raise ValueError(f"API error: {data['error']['message']}")
    return data.get('items', [])

def show_yt_selector(app_data, items, original_query: str):
    """Popup to select videos with checkboxes, thumbnails, and action buttons."""
    root = tk.Toplevel()
    root.title("YouTube Suggestions")
    root.geometry("500x400")
    root.resizable(False, False)

    # Import PIL for image handling
    try:
        from PIL import Image, ImageTk
        import io
        PIL_AVAILABLE = True
    except ImportError:
        PIL_AVAILABLE = False
        print("[INFO] Consider installing Pillow for thumbnail support: pip install Pillow")

    vars_ = []
    video_map = {}

    # Main container
    main_frame = Frame(root)
    main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

    # Scrollable frame for video items
    canvas = Canvas(main_frame)
    scrollbar = Scrollbar(main_frame, orient="vertical", command=canvas.yview)
    scrollable_frame = Frame(canvas)

    scrollable_frame.bind(
        "<Configure>",
        lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
    )

    canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
    canvas.configure(yscrollcommand=scrollbar.set)

    for i, item in enumerate(items):
        # Get video data
        snippet = item['snippet']
        video_id = item['id']['videoId']
        title = snippet['title']
        url = f"https://www.youtube.com/watch?v={video_id}"

        # Get thumbnail URL (using medium quality)
        thumbnails = snippet.get('thumbnails', {})
        thumb_url = thumbnails.get('medium', {}).get('url') or \
                   thumbnails.get('high', {}).get('url') or \
                   thumbnails.get('default', {}).get('url')

        # Create frame for each video item
        item_frame = Frame(scrollable_frame, relief="groove", borderwidth=1)
        item_frame.pack(fill=tk.X, pady=5, padx=5)

        if PIL_AVAILABLE and thumb_url:
            try:
                # Fetch and display thumbnail
                response = requests.get(thumb_url, timeout=5)
                img_data = Image.open(io.BytesIO(response.content))
                img_data = img_data.resize((120, 90), Image.Resampling.LANCZOS)
                photo = ImageTk.PhotoImage(img_data)

                # Keep a reference to prevent garbage collection
                if not hasattr(show_yt_selector, 'images'):
                    show_yt_selector.images = []
                show_yt_selector.images.append(photo)

                # Thumbnail label
                thumb_label = Label(item_frame, image=photo)
                thumb_label.pack(side=tk.LEFT, padx=5, pady=5)
            except Exception as e:
                print(f"[WARNING] Could not load thumbnail: {e}")
                # Fallback to text-only if image loading fails
                thumb_label = Label(item_frame, text="[No Thumbnail]", width=15)
                thumb_label.pack(side=tk.LEFT, padx=5, pady=5)
        else:
            # Placeholder when PIL is not available
            thumb_label = Label(item_frame, text="[Thumbnail]", width=15, bg="lightgray")
            thumb_label.pack(side=tk.LEFT, padx=5, pady=5)

        # Title and checkbox frame
        text_frame = Frame(item_frame)
        text_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=5, pady=5)

        # Video title (truncated if too long)
        display_title = title[:60] + "..." if len(title) > 60 else title
        title_label = Label(text_frame, text=display_title, wraplength=250, justify="left")
        title_label.pack(anchor="w")

        # Checkbox
        var = IntVar()
        vars_.append(var)
        check_btn = Checkbutton(text_frame, text="Select", variable=var)
        check_btn.pack(anchor="w")

        # Store video data
        video_map[title] = {
            'title': title,
            'url': url,
            'thumbnail_url': thumb_url
        }

    # Pack scrollbar and canvas
    scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
    canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

    def add_selected_to_doc():
        selected_titles = [items[i]['snippet']['title'] for i, var in enumerate(vars_) if var.get()]
        if not selected_titles:
            messagebox.showwarning("No Selection", "No videos selected!")
            return
        selected_urls = [video_map[title]['url'] for title in selected_titles]
        insert_yt_to_docx(app_data, selected_urls, original_query, all=False)
        root.destroy()

    def add_all_to_doc():
        urls = [video_map[items[i]['snippet']['title']]['url'] for i in range(len(items))]
        insert_yt_to_docx(app_data, urls, original_query, all=True)
        root.destroy()

    # Button frame
    button_frame = Frame(root)
    button_frame.pack(fill=tk.X, pady=10)

    Button(button_frame, text="ASTD (Add Selected To Doc)", command=add_selected_to_doc).pack(side=tk.LEFT, padx=20, pady=5)
    Button(button_frame, text="AATD (Add All To Doc)", command=add_all_to_doc).pack(side=tk.RIGHT, padx=20, pady=5)

    # root.mainloop()
    root.wait_window() # using this as thr root is obj with Toplevel not Tk

def insert_yt_to_docx(app_data, urls: list, original_query: str, all: bool):
    """Add selected/all YT links as bullets in the current DOCX."""
    path = app_data.get_folder_path()  + "/" + app_data.get_file_name()
    if not path:
        raise ValueError("No DOCX open!")
    doc = Document(path)
    p = doc.add_paragraph(f"YouTube Suggestions for '{original_query}':", style="Heading 2")
    p.paragraph_format.space_after = 0
    for url in urls:
        p = doc.add_paragraph("", style="List Bullet")
        # add_hyperlink(p, url, "Watch Video")
        add_hyperlink(p, url, url)
        p.paragraph_format.left_indent = Inches(0.5)
    doc.save(path)
    messagebox.showinfo("Success", f"Added {'all' if all else 'selected'} YT suggestions to notes!")

def handle_yt_from_text(app_data, text):
    """Main handler: Fetch and show suggestions."""
    if not text.strip():
        messagebox.showwarning("No Text", "Select text first!")
        return
    try:
        items = fetch_yt_suggestions(text)
        if not items:
            messagebox.showinfo("No Results", "No YT suggestions found.")
            return
        show_yt_selector(app_data, items, text)
    except Exception as e:
        messagebox.showerror("YT Error", f"Failed to fetch suggestions: {e}")

# if __name__ == "__main__":
#     handle_yt_from_text("Python tutorials")